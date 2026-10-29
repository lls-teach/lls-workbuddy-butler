#!/usr/bin/env python3
"""从大纲生成 PPT v2.2（版式词汇：KPI 行 / 语义状态表 / 进度条 / 斑马表 / 来源行）。

用法: python3 pptx_create.py -o 成片.pptx --title 标题 --subtitle 副标题 --outline 大纲.md [--style formal]
大纲语法:
  '# ' 页标题（写结论不写话题）
  '[KPI] 31,940|技能总量|仍在高速增长'        KPI 数字行
  '[STATUS] 发票报销管家|163 竞品仅 3.5 万次|11 月中旬|研发中|amber'
  '[BAR] 58|实名认证'                          进度条
  '[TABLE] 合作项|投入|周期'                   斑马表（首行表头）
  '- ' 要点（>=5 两列）   '> ' 结论带   '= ' 来源行
"""
import argparse
import math
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from lls_pptx_utils import apply_cn_font, CN_FONT, die, ensure_parent, guard_overwrite, parse_outline

STYLES = {
    "formal": {
        "bg": RGBColor(0xFA, 0xFA, 0xF7), "fg": RGBColor(0x21, 0x24, 0x28),
        "muted": RGBColor(0x6E, 0x74, 0x7B), "accent": RGBColor(0x1F, 0x46, 0x9C),
        "rule": RGBColor(0xD9, 0xDC, 0xE0), "zebra": RGBColor(0xF1, 0xEF, 0xE9),
        "title_pt": 26, "body_pt": 17, "cover_title_pt": 40,
    },
    "consulting": {
        "bg": RGBColor(0xFF, 0xFF, 0xFF), "fg": RGBColor(0x2B, 0x2B, 0x2B),
        "muted": RGBColor(0x77, 0x77, 0x77), "accent": RGBColor(0x8C, 0x1F, 0x3F),
        "rule": RGBColor(0xE0, 0xE0, 0xE0), "zebra": RGBColor(0xF2, 0xF2, 0xF2),
        "title_pt": 24, "body_pt": 15, "cover_title_pt": 38,
    },
    "teaching": {
        "bg": RGBColor(0xFD, 0xFB, 0xF3), "fg": RGBColor(0x30, 0x33, 0x38),
        "muted": RGBColor(0x7A, 0x7F, 0x85), "accent": RGBColor(0xE2, 0x6D, 0x2C),
        "rule": RGBColor(0xEA, 0xE4, 0xD5), "zebra": RGBColor(0xF5, 0xF0, 0xE2),
        "title_pt": 28, "body_pt": 20, "cover_title_pt": 44,
    },
}
# 语义色：状态关键词 → 颜色；第 5 列可显式指定 amber/blue/good/muted
SEMANTIC = {
    "amber": RGBColor(0xB7, 0x79, 0x1F), "blue": RGBColor(0x2B, 0x6C, 0xB0),
    "good": RGBColor(0x2F, 0x85, 0x5A), "muted": RGBColor(0x8A, 0x8F, 0x96),
}
STATUS_HINT = [("完成", "good"), ("上线", "good"), ("通过", "good"), ("研发", "amber"),
               ("进行", "amber"), ("内测", "amber"), ("冻结", "blue"), ("设计", "blue"),
               ("排队", "muted"), ("计划", "muted")]
BANDS = {"title_top": 0.42, "title_h": 0.85, "content_top": 1.55,
         "content_bottom": 6.75, "source": 6.42, "footer": 7.05}


def sem_color(status: str, S):
    for key, color in SEMANTIC.items():
        if status == key:
            return color
    for word, key in STATUS_HINT:
        if word in status:
            return SEMANTIC[key]
    return S["accent"]


def new_slide(prs, S):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = S["bg"]
    return slide


def add_text(slide, text, left, top, width, height, size, color, bold=False):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    for run in p.runs:
        apply_cn_font(run)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def add_rule(slide, left, top, width, color, thick=0.028):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(thick))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_bullets(slide, bullets, S, left, top, width, height):
    if not bullets:
        return 0
    two_col = len(bullets) >= 5
    cols = [bullets[:math.ceil(len(bullets) / 2)], bullets[math.ceil(len(bullets) / 2):]] if two_col else [bullets]
    col_w = (width - 0.5) / 2 if two_col else width
    for ci, items in enumerate(cols):
        if not items:
            continue
        box = slide.shapes.add_textbox(Inches(left + ci * (col_w + 0.5) if two_col else left),
                                       Inches(top), Inches(col_w), Inches(height))
        tf = box.text_frame
        tf.word_wrap = True
        for j, b in enumerate(items):
            para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            para.text = "• " + b
            for run in para.runs:
                apply_cn_font(run)
                run.font.size = Pt(S["body_pt"])
                run.font.color.rgb = S["fg"]
            para.space_after = Pt(12)
    return 1


def add_kpis(slide, kpis, S, top):
    n = min(len(kpis), 4)
    col_w = 11.8 / n
    for i, k in enumerate(kpis[:n]):
        num, label, note = (k + ["", ""])[:3]
        x = 0.75 + i * col_w
        add_text(slide, num, x, top, col_w - 0.3, 0.85, 36, S["fg"], bold=True)
        add_text(slide, label, x, top + 0.85, col_w - 0.3, 0.35, 13, S["fg"])
        if note:
            add_text(slide, note, x, top + 1.2, col_w - 0.3, 0.35, 11, S["muted"])


def add_status_table(slide, rows, S, top, height):
    row_h = min(0.72, height / max(len(rows), 1))
    for i, row in enumerate(rows):
        name, gap, when, status = (row + [""] * 4)[:4]
        color_key = row[4] if len(row) > 4 and row[4] in SEMANTIC else None
        color = SEMANTIC[color_key] if color_key else sem_color(status, S)
        y = top + i * row_h
        if i % 2 == 0:
            z = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(y - 0.06),
                                       Inches(11.9), Inches(row_h))
            z.fill.solid(); z.fill.fore_color.rgb = S["zebra"]
            z.line.fill.background(); z.shadow.inherit = False
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.85), Inches(y + 0.14), Inches(0.16), Inches(0.16))
        dot.fill.solid(); dot.fill.fore_color.rgb = color
        dot.line.fill.background(); dot.shadow.inherit = False
        add_text(slide, name, 1.2, y, 3.2, row_h, 15, S["fg"], bold=True)
        if gap:
            add_text(slide, gap, 4.5, y, 4.6, row_h, 13, S["fg"])
        if when:
            add_text(slide, when, 9.2, y, 1.4, row_h, 13, S["muted"])
        if status:
            add_text(slide, status, 10.7, y, 1.8, row_h, 13, color)


def add_bars(slide, bars, S, top, height):
    row_h = min(0.6, height / max(len(bars), 1))
    for i, b in enumerate(bars):
        pct, label = (b + ["", ""])[:2]
        y = top + i * row_h
        try:
            v = max(0.0, min(100.0, float(pct)))
        except ValueError:
            continue
        track = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3.6), Inches(y + 0.09),
                                       Inches(7.0), Inches(0.26))
        track.fill.solid(); track.fill.fore_color.rgb = S["rule"]
        track.line.fill.background(); track.shadow.inherit = False
        fillw = max(0.05, 7.0 * v / 100.0)
        fill = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3.6), Inches(y + 0.09),
                                      Inches(fillw), Inches(0.26))
        fill.fill.solid(); fill.fill.fore_color.rgb = S["accent"]
        fill.line.fill.background(); fill.shadow.inherit = False
        add_text(slide, label, 0.75, y, 2.7, row_h, 14, S["fg"])
        add_text(slide, f"{v:g}%", 10.75, y, 1.5, row_h, 14, S["fg"], bold=True)


def add_table(slide, rows, S, top, height):
    row_h = min(0.6, height / max(len(rows), 1))
    ncols = max(len(r) for r in rows)
    col_w = 11.9 / ncols
    for i, row in enumerate(rows):
        y = top + i * row_h
        if i % 2 == 1 and i > 0:
            z = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(y - 0.05),
                                       Inches(11.9), Inches(row_h))
            z.fill.solid(); z.fill.fore_color.rgb = S["zebra"]
            z.line.fill.background(); z.shadow.inherit = False
        for j in range(ncols):
            text = row[j] if j < len(row) else ""
            color = S["fg"] if i == 0 or j == 0 else (S["accent"] if j == 1 and ncols >= 3 else S["fg"])
            add_text(slide, text, 0.75 + j * col_w, y, col_w - 0.2, row_h, 13.5, color, bold=(i == 0 or j == 0))


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--output", required=True)
    ap.add_argument("--title", required=True)
    ap.add_argument("--subtitle", default="")
    ap.add_argument("--outline", required=True)
    ap.add_argument("--style", choices=list(STYLES), default="formal")
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()

    S = STYLES[args.style]
    # utf-8-sig：兼容 Windows 记事本保存的带 BOM 大纲（否则首行 "# " 会因
    # BOM 前缀识别失败、被当成要点）；无 BOM 时与 utf-8 完全等价。
    with open(args.outline, encoding="utf-8-sig") as f:
        outline = parse_outline(f.read())
    if not outline:
        die("大纲为空或无法解析")
    total = len(outline) + 1

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    cover = new_slide(prs, S)
    add_text(cover, args.title, 1.0, 2.7, 11.3, 1.4, S["cover_title_pt"], S["fg"], bold=True)
    if args.subtitle:
        add_text(cover, args.subtitle, 1.0, 4.3, 11.3, 0.7, 18, S["muted"])
    add_rule(cover, 1.0, 4.05, 2.2, S["accent"])

    for i, sd in enumerate(outline, 1):
        slide = new_slide(prs, S)
        add_text(slide, args.title, 0.6, 0.16, 9.0, 0.32, 10, S["muted"])
        title_text = f"{i}. {sd['title']}" if sd["title"] else str(i)
        add_text(slide, title_text, 0.6, BANDS["title_top"], 12.1, BANDS["title_h"],
                 S["title_pt"], S["fg"], bold=True)
        add_rule(slide, 0.6, 1.42, 12.1, S["rule"])

        content_h = BANDS["content_bottom"] - BANDS["content_top"] - (0.55 if sd.get("takeaway") else 0)
        cursor = BANDS["content_top"]
        if sd["kpis"]:
            add_kpis(slide, sd["kpis"], S, cursor)
            cursor += 1.75
        if sd["bullets"]:
            add_bullets(slide, sd["bullets"], S, 0.75, cursor, 11.8,
                        max(1.0, content_h - (cursor - BANDS["content_top"])))
            cursor += min(len(sd["bullets"]) * 0.55, 2.6)
        if sd["bars"]:
            add_bars(slide, sd["bars"], S, cursor, max(1.0, BANDS["content_bottom"] - 0.7 - cursor))
            cursor += min(len(sd["bars"]) * 0.6, 3.0)
        bottom = BANDS["content_bottom"] - (0.55 if sd.get("takeaway") else 0)
        if sd["status_rows"]:
            add_status_table(slide, sd["status_rows"], S, cursor, max(1.0, bottom - cursor))
        if sd["table_rows"]:
            add_table(slide, sd["table_rows"], S, cursor, max(1.0, bottom - cursor))
        if sd.get("takeaway"):
            takeaway = re.sub(r"^结论[:：]\s*", "", sd["takeaway"])
            tk_y = BANDS["content_bottom"] - 0.45 - (0.42 if sd.get("source") else 0)
            add_text(slide, "本页结论：" + takeaway, 0.75, tk_y, 11.8, 0.5,
                     S["body_pt"] - 2, S["accent"], bold=True)
        if sd.get("source"):
            # 大纲写法本就是 "= 来源：…"，先剥掉自带前缀，避免 "来源：来源："
            src_text = re.sub(r"^来源[:：]\s*", "", sd["source"])
            add_text(slide, "来源：" + src_text, 0.75, BANDS["footer"] - 0.22, 11.8, 0.3, 10, S["muted"])
        folio = slide.shapes.add_textbox(Inches(12.3), Inches(BANDS["footer"]), Inches(0.8), Inches(0.35))
        fp = folio.text_frame.paragraphs[0]
        fp.text = f"{i + 1} / {total}"
        for run in fp.runs:
            apply_cn_font(run)
            run.font.size = Pt(10)
            run.font.color.rgb = S["muted"]

    out = guard_overwrite(args.output, args.force)
    ensure_parent(out)
    prs.save(out)
    kinds = sum(1 for sd in outline if sd["kpis"] or sd["status_rows"] or sd["bars"] or sd["table_rows"])
    print(f"[lls-pptx] 已生成 {total} 页（风格 {args.style}，结构化页 {kinds}）→ {out}")
    print(f"[lls-pptx] 下一步：$PY scripts/pptx_check.py {out}")


if __name__ == "__main__":
    main()
