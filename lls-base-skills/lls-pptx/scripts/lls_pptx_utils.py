"""lls-pptx 共享工具：覆盖保护、中文字体设置、大纲解析、单行 token 保护、图片适配。"""
from __future__ import annotations

import os
import re
import sys

from pptx.oxml.ns import qn

CN_FONT = "微软雅黑"


def die(msg: str, code: int = 1) -> None:
    print(f"[lls-pptx] 错误: {msg}", file=sys.stderr)
    sys.exit(code)


def ensure_parent(path: str) -> str:
    parent = os.path.dirname(os.path.abspath(path))
    os.makedirs(parent, exist_ok=True)
    return path


def guard_overwrite(path: str, force: bool = False) -> str:
    if not os.path.exists(path) or force:
        return path
    root, ext = os.path.splitext(path)
    alt = f"{root}_new{ext}"
    print(f"[lls-pptx] 目标已存在，改写到 {alt}（要覆盖请加 --force）", file=sys.stderr)
    return alt


def parse_outline(text: str) -> list[dict]:
    """v3 大纲解析（版式词汇）：

      '# '  页标题（写结论不写话题）
      '[KPI] 31,940|技能总量|仍在高速增长'        KPI 数字行（一页最多 4 个）
      '[STATUS] 名称|空位说明|时间|状态|色'        语义状态表行（色: amber/blue/good/muted，缺省按状态关键词）
      '[BAR] 58|实名认证'                          进度条行（百分比）
      '[TABLE] 合作项|我方投入|贵方投入|周期'       斑马纹表格行（首个 TABLE 行为表头）
      '[IMAGE] 图片路径|说明|contain'              图片页（contain 等比居中 / cover 先裁后放，缺省 contain）
      '- '  要点（>=5 自动两列）
      '> '  本页结论（结论带）
      '= '  来源行（页脚上方）
    """
    # 防御性剥离 BOM：大纲若经其他路径读入而未用 utf-8-sig，首行仍可识别。
    text = text.removeprefix("\ufeff")
    slides: list[dict] = []

    def new_page():
        return {"title": "", "kpis": [], "bars": [], "status_rows": [], "table_rows": [],
                "bullets": [], "takeaway": None, "source": None, "images": []}

    cur = None
    for raw in text.split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("# "):
            cur = new_page(); cur["title"] = line[2:].strip(); slides.append(cur)
            continue
        if cur is None:
            cur = new_page(); slides.append(cur)
        if line.startswith("[KPI] "):
            parts = [p.strip() for p in line[6:].split("|")]
            cur["kpis"].append(parts + [""] * (3 - len(parts)))
        elif line.startswith("[STATUS] "):
            parts = [p.strip() for p in line[9:].split("|")]
            cur["status_rows"].append(parts + [""] * (5 - len(parts)))
        elif line.startswith("[BAR] "):
            parts = [p.strip() for p in line[6:].split("|")]
            cur["bars"].append(parts + [""] * (2 - len(parts)))
        elif line.startswith("[TABLE] "):
            cur["table_rows"].append([p.strip() for p in line[8:].split("|")])
        elif line.startswith("[IMAGE] "):
            parts = [p.strip() for p in line[8:].split("|")]
            cur["images"].append({"path": parts[0] if parts else "",
                                 "caption": parts[1] if len(parts) > 1 else "",
                                 "mode": parts[2] if len(parts) > 2 else "contain"})
        elif line.startswith("> "):
            cur["takeaway"] = line[2:].strip()
        elif line.startswith("= "):
            cur["source"] = line[2:].strip()
        elif line.startswith("- "):
            cur["bullets"].append(line[2:].strip())
        else:
            cur["bullets"].append(line.strip())
    return [s for s in slides if s["title"] or s["bullets"] or s["kpis"] or s["status_rows"] or s["bars"] or s["table_rows"] or s["images"]]


def apply_cn_font(run, name: str = CN_FONT) -> None:
    """python-pptx 的 font.name 只写 a:latin；CJK 必须同时写 a:ea，否则部分渲染器丢字形。"""
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", name)


# ---- 单行 token 保护：页码、KPI 数字、百分比、金额、日期、短代码窄盒防换行堆叠 ----

LATIN_EM = 0.62      # 拉丁字母/数字约 0.62em
CJK_EM = 1.0         # 全角字符约 1.0em（原 0.62 统一系数对中文偏小，此处按字符区分）
TOKEN_PADDING_IN = 0.14
TOKEN_SAFETY = 1.15


def _char_em(ch: str) -> float:
    o = ord(ch)
    if (0x4E00 <= o <= 0x9FFF or 0x3400 <= o <= 0x4DBF or 0xF900 <= o <= 0xFAFF
            or 0x3000 <= o <= 0x303F or 0xFF00 <= o <= 0xFFEF):
        return CJK_EM
    return LATIN_EM


def estimate_single_line_width(text, font_size_pt, padding_in=TOKEN_PADDING_IN, safety=TOKEN_SAFETY) -> float:
    """估算单行文本所需宽度（英寸）：字符 em 之和 × 字号 / 72 + 边距，再 × 安全系数。"""
    total = sum(_char_em(ch) for ch in str(text).strip())
    return (total * float(font_size_pt) / 72.0 + padding_in) * safety


_SHORT_TOKEN_RES = [
    re.compile(r"^\d{1,4}$"),
    re.compile(r"^\d+(?:[.,]\d+)?%$"),
    re.compile(r"^[¥$€£]\s?\d+(?:[.,]\d+)?(?:[KMB万亿])?$", re.IGNORECASE),
    re.compile(r"^\d{4}(?:[-/.]\d{1,2}){0,2}$"),
    re.compile(r"^\d{4}\s?Q[1-4]$", re.IGNORECASE),
    re.compile(r"^(?:W|P|Q|H)\d{1,3}$", re.IGNORECASE),
    re.compile(r"^[A-Z0-9][A-Z0-9+%./:–—-]{0,11}$"),
]


def is_short_token(text) -> bool:
    """识别短代码：纯数字/百分比/金额/日期/季度/周次/大写代码，≤16 字、无双空格、无换行。"""
    t = str(text or "").strip()
    if not t or len(t) > 16 or re.search(r"\s{2,}|[\r\n]", t):
        return False
    return any(rx.match(t) for rx in _SHORT_TOKEN_RES)


def add_single_line_token(slide, text, left, top, width, height, size, color, bold=False,
                          min_size=12, font_name=CN_FONT):
    """单行 token 文本框：不换行、零边距、字号自动缩小到下限、仍窄则告警。

    python-pptx 侧等价物：word_wrap=False + margins 0 + TEXT_TO_FIT_SHAPE（缩小兜底）。
    语义已知的 token（页码、KPI 数字）直接走本函数，不必先过 is_short_token。
    """
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Inches, Pt
    requested = float(size)
    safe_min = min(requested, max(1.0, float(min_size or 12)))
    recommended = estimate_single_line_width(text, requested)
    effective = requested
    if width < recommended:
        effective = max(safe_min, requested * width / recommended)
    minimum = estimate_single_line_width(text, effective)
    if width < minimum:
        print(f"[lls-pptx] 单行 token 可能仍太窄：{str(text)!r} 框 {width:.2f}in，"
              f"需 {minimum:.2f}in（{effective:.1f}pt）", file=sys.stderr)
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = str(text)
    for run in p.runs:
        apply_cn_font(run, font_name)
        run.font.size = Pt(effective)
        run.font.bold = bold
        run.font.color.rgb = color
    return {"adjusted": effective < requested, "requested_font_size": requested,
            "font_size": effective, "recommended_width": recommended, "box_width": width}


# ---- 图片适配：等比 contain / 先裁后放 cover ----

def image_size(path):
    """返回图片像素尺寸 (w, h)。"""
    from PIL import Image
    with Image.open(path) as im:
        return im.size


def contain_rect(img_w, img_h, box_x, box_y, box_w, box_h):
    """等比放进目标框（不变形），返回居中后的 (x, y, w, h)，单位与 box 一致。"""
    s = min(box_w / img_w, box_h / img_h)
    w, h = img_w * s, img_h * s
    return (box_x + (box_w - w) / 2, box_y + (box_h - h) / 2, w, h)


def cover_image_bytes(path, box_w_in, box_h_in, dpi=200, position="center"):
    """cover：按目标框比例先裁后放，返回 (BytesIO PNG, 目标像素宽, 目标像素高）。

    先裁好再放，避免渲染器对裁剪比例解释不一致导致走样；position: center/north/south。
    """
    from io import BytesIO
    from PIL import Image, ImageOps
    tw, th = max(1, round(box_w_in * dpi)), max(1, round(box_h_in * dpi))
    with Image.open(path) as im:
        im = ImageOps.exif_transpose(im)
        iw, ih = im.size
        scale = max(tw / iw, th / ih)
        nw, nh = max(1, round(iw * scale)), max(1, round(ih * scale))
        im = im.resize((nw, nh), Image.LANCZOS)
        if position == "north":
            top = 0
        elif position == "south":
            top = nh - th
        else:
            top = (nh - th) // 2
        left = (nw - tw) // 2
        im = im.crop((left, top, left + tw, top + th))
        bio = BytesIO()
        im.save(bio, format="PNG")
        bio.seek(0)
        return bio, tw, th
