#!/usr/bin/env python3
"""程序化检查 PPT：密度、要点数、空占位、中文字体、页数健康度、单行 token 溢出、图片拉伸。零渲染、秒级。

用法: python3 pptx_check.py 输入.pptx [--max-bullets 6] [--max-chars 600]
退出码: 0=通过 1=有警告（不阻塞，供 AI 汇报）
"""
import argparse
import re
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from lls_pptx_utils import CN_FONT, die, estimate_single_line_width

warnings: list[str] = []


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("--max-bullets", type=int, default=6)
    ap.add_argument("--max-chars", type=int, default=600, help="单页正文文本上限（密度带宽启发值）")
    args = ap.parse_args()

    try:
        prs = Presentation(args.input)
    except Exception as e:  # noqa: BLE001
        die(f"无法打开: {e}")

    n = len(prs.slides)
    print(f"[lls-pptx] 检查 {args.input}（{n} 页）")

    if n == 0:
        warnings.append("空演示文稿")
    if n > 40:
        warnings.append(f"页数过多（{n} 页）——超过 40 页建议拆分成套")

    for i, slide in enumerate(prs.slides, 1):
        texts, bullet_count, char_count, fonts = [], 0, 0, set()
        title_text = None
        empty_frames = 0
        text_shapes = 0
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            # 只把真文本框的空壳算"空占位符"；矩形分隔线等 AUTO_SHAPE 天带空 text_frame
            if not shape.text_frame.text.strip():
                if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
                    empty_frames += 1
                continue
            text_shapes += 1
            # 标题候选：首个字号 > 12pt 的文本（跳过 10pt 页眉运行标题，否则
            # 每页都把页眉误判为标题）；都没有则回退首个文本。
            if title_text is None:
                sizes = [r.font.size.pt for p in shape.text_frame.paragraphs
                         for r in p.runs if r.font.size]
                if not sizes or max(sizes) > 12:
                    title_text = shape.text_frame.text.strip().split("\n")[0]
            for para in shape.text_frame.paragraphs:
                t = "".join(r.text for r in para.runs).strip()
                if not t:
                    continue
                texts.append(t)
                if t.startswith(("•", "-", "*")):
                    bullet_count += 1
                char_count += len(t)
                for r in para.runs:
                    if r.font.name:
                        fonts.add(r.font.name)
        if bullet_count > args.max_bullets:
            warnings.append(f"第 {i} 页要点 {bullet_count} 条 > {args.max_bullets}——建议拆页或转两列证据网格")
        if char_count > args.max_chars:
            warnings.append(f"第 {i} 页正文 {char_count} 字 > {args.max_chars}——密度超标，检查是否该拆分")
        if empty_frames:
            warnings.append(f"第 {i} 页有 {empty_frames} 个空文本框（空占位符）")
        if texts:
            title = title_text or texts[0]
            topic_markers = ("分析", "介绍", "概述", "关于", "情况", "简介")
            # 封面页的标题本就是演示主题（如"数据看板入门"），不该按内容页要求改成结论句。
            # 判定封面用"第 1 页 + 文本块少 + 总字数少"，避免把开门就上内容的首屏误当封面放过。
            is_cover = i == 1 and text_shapes <= 3 and char_count < 60
            looks_topic = (
                any(m in title for m in topic_markers)
                or re.fullmatch(r"[\d一二三四五六七八九十]+[.、]", title.strip()) is not None
                or (
                    len(title.strip()) <= 8
                    and not any(ch in title for ch in "：:，,。；")
                )
            )
            if looks_topic and not is_cover:
                warnings.append(f"第 {i} 页标题疑似话题式（{title!r}）——建议写结论句")
        cjk_runs = [f for f in fonts if f and f != CN_FONT and any('\u4e00' <= ch <= '\u9fff' for ch in "".join(texts))]
        if CN_FONT not in fonts and texts:
            warnings.append(f"第 {i} 页未设置 {CN_FONT} 字体（渲染环境可能缺字）")
        # 单行 token 溢出：word_wrap 关掉的框按单行估算（先按请求字号算，缩到 8pt 仍塞不下才报）
        for shape in slide.shapes:
            if not shape.has_text_frame or shape.text_frame.word_wrap is not False:
                continue
            for para in shape.text_frame.paragraphs:
                t = "".join(r.text for r in para.runs).strip()
                if not t:
                    continue
                sizes = [r.font.size.pt for r in para.runs if r.font.size]
                size = max(sizes) if sizes else 18
                est = estimate_single_line_width(t, size)
                if est <= shape.width.inches:
                    continue
                eff = max(8.0, size * shape.width.inches / est)
                if estimate_single_line_width(t, eff) > shape.width.inches:
                    warnings.append(f"第 {i} 页单行文本可能溢出（{t[:12]!r} 需 {est:.2f}in，框 {shape.width.inches:.2f}in）——加宽盒子或缩小字号")
        # 图片拉伸：框比例与原图比例差 >2% 即报（比例比较与单位无关）
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            try:
                iw, ih = shape.image.size
                fw, fh = shape.width.inches, shape.height.inches
                if iw and ih and fw and fh and abs(iw / ih - fw / fh) / (iw / ih) > 0.02:
                    warnings.append(f"第 {i} 页图片疑似被拉伸（原图 {iw}×{ih}，框 {fw:.2f}×{fh:.2f}in）——应走 contain/cover")
            except Exception:  # noqa: BLE001
                pass

    # 汇总
    if warnings:
        print(f"[lls-pptx] ⚠️ {len(warnings)} 条警告:")
        for w in warnings:
            print("  - " + w)
        sys.exit(1)
    print("[lls-pptx] ✅ 全部通过：无超密度页、无空占位、要点数健康")


if __name__ == "__main__":
    main()
