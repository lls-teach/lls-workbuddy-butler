#!/usr/bin/env python3
"""读取 PPT 全部文本（逐页列出），可选输出 JSON。

用法: python3 pptx_read.py 输入.pptx [--json out.json]
"""
import argparse
import json

from pptx import Presentation

from lls_pptx_utils import die, ensure_parent


def slide_text(slide) -> list[str]:
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                t = "".join(run.text for run in para.runs).strip()
                if t:
                    texts.append(t)
        if shape.has_table:
            for row in shape.table.rows:
                cells = [c.text.strip() for c in row.cells]
                texts.append(" | ".join(cells))
    return texts


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("--json", help="输出 JSON 文件")
    args = ap.parse_args()

    try:
        prs = Presentation(args.input)
    except Exception as e:  # noqa: BLE001
        die(f"无法打开（.ppt 旧格式请先另存为 .pptx）: {e}")

    pages = []
    for i, slide in enumerate(prs.slides, 1):
        texts = slide_text(slide)
        pages.append({"slide": i, "texts": texts})
        print(f"===== 第 {i} 页 =====")
        for t in texts:
            print("  " + t[:90])

    print(f"[lls-pptx] 共 {len(pages)} 页")
    if args.json:
        ensure_parent(args.json)
        # with 块关闭句柄：Windows 下未关闭的文件句柄会锁文件、影响后续读写。
        with open(args.json, "w", encoding="utf-8") as f:
            json.dump(pages, f, ensure_ascii=False, indent=1)
        print(f"[lls-pptx] JSON → {args.json}")


if __name__ == "__main__":
    main()
