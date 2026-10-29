#!/usr/bin/env python3
"""教学设计 JSON → 教案 docx（要素齐全）。"""
import argparse
import json

from docx import Document
from docx.shared import Pt

from lls_jiaoan_utils import ensure_parent, guard_overwrite, set_cn_font


def para(doc, text, size=12, bold=False, align=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    set_cn_font(run)
    if align is not None:
        p.alignment = align
    return p


def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(it)
        run.font.size = Pt(12)
        set_cn_font(run)


def section(doc, title):
    para(doc, title, size=14, bold=True)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("json")
    ap.add_argument("-o", "--output", required=True)
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()

    from lls_jiaoan_utils import die
    try:
        with open(args.json, encoding="utf-8") as f:
            e = json.load(f)
    except (OSError, ValueError) as exc:
        die(f"无法读取教学设计 JSON: {exc}")
    if not isinstance(e, dict):
        die("教学设计 JSON 顶层须为对象")
    for k in ("topic", "grade", "subject", "goals", "key_points", "difficult_points", "process"):
        if not e.get(k):
            die(f"要素缺失: {k}")
    if not isinstance(e["goals"], dict):
        die("要素格式错误: goals 须为对象")
    if not isinstance(e["process"], list) or not all(isinstance(s, dict) for s in e["process"]):
        die("要素格式错误: process 须为对象数组")
    for k in ("key_points", "difficult_points"):
        if not isinstance(e[k], list):
            die(f"要素格式错误: {k} 须为数组")
    if e.get("homework") is not None and not isinstance(e["homework"], list):
        die("要素格式错误: homework 须为数组")
    if e.get("prep") is not None and not isinstance(e["prep"], dict):
        die("要素格式错误: prep 须为对象")
    for key in ("knowledge", "skill", "attitude"):
        if e["goals"].get(key) is not None and not isinstance(e["goals"][key], list):
            die(f"要素格式错误: goals.{key} 须为数组")

    doc = Document()
    para(doc, e["topic"], size=16, bold=True, align=1)
    para(doc, f"{e.get('grade','')} · {e.get('subject','')} · {e.get('duration','45 分钟')}", size=11, align=1)

    section(doc, "一、教学目标")
    g = e["goals"]
    for label, key in [("知识目标", "knowledge"), ("能力目标", "skill"), ("情感态度", "attitude")]:
        if g.get(key):
            para(doc, f"{label}：", size=12, bold=True)
            bullets(doc, g[key])

    section(doc, "二、教学重难点")
    para(doc, "重点：", size=12, bold=True)
    bullets(doc, e["key_points"])
    para(doc, "难点：", size=12, bold=True)
    bullets(doc, e["difficult_points"])

    section(doc, "三、教学准备")
    prep = e.get("prep") or {}
    para(doc, f"教师：{'、'.join(prep.get('teacher', ['课件']))}")
    para(doc, f"学生：{'、'.join(prep.get('student', ['课本、练习本']))}")

    section(doc, "四、教学过程")
    for i, step in enumerate(e["process"], 1):
        para(doc, f"环节 {i}：{step.get('phase','')}（{step.get('time','')}）", size=12, bold=True)
        para(doc, f"教师活动：{step.get('teacher','')}")
        para(doc, f"学生活动：{step.get('student','')}")
        if step.get("design"):
            para(doc, f"设计意图：{step.get('design')}")
        para(doc, "", size=6)

    section(doc, "五、板书设计")
    para(doc, e.get("board", "（板书由课件呈现）"))

    section(doc, "六、作业布置")
    bullets(doc, e.get("homework", ["完成课后练习"]))

    section(doc, "七、教学反思")
    para(doc, e.get("reflection_hint", "（课后填写）"))

    out = guard_overwrite(args.output, args.force)
    ensure_parent(out)
    doc.save(out)
    print(f"[lls-jiaoan] 教案已生成 → {out}（环节 {len(e['process'])} 个）")


main()
