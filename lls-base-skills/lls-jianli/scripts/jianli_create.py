#!/usr/bin/env python3
"""从结构化 JSON 生成专业单栏中文简历（python-docx，宋体/雅黑，教育-经历-项目-技能四模块）。

用法: python3 jianli_create.py -o 简历.docx --data 简历.json
JSON 结构见 SKILL.md「简历信息清单」。
"""
import argparse
import json

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from lls_jianli_utils import die, ensure_parent, guard_overwrite, set_cn_font

REQUIRED = ["name", "phone", "email"]


def para(doc, text, size=10.5, bold=False, align=None, space_after=4, color=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    set_cn_font(run)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def section(doc, title):
    para(doc, title, size=13, bold=True, space_after=6)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--output", required=True)
    ap.add_argument("--data", required=True, help="简历 JSON")
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()

    data = json.load(open(args.data, encoding="utf-8"))
    missing = [k for k in REQUIRED if not data.get(k)]
    if missing:
        die(f"简历 JSON 缺必填字段: {missing}（name/phone/email 是简历最低配置）")

    doc = Document()
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Pt(36)
        sec.left_margin = sec.right_margin = Pt(46)

    para(doc, data["name"], size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    contact = " | ".join(filter(None, [data.get("phone"), data.get("email"), data.get("city", "")]))
    para(doc, contact, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10, color="555555")
    if data.get("summary"):
        para(doc, data["summary"], size=11, bold=True, space_after=10)

    for key, title in [("education", "教育背景"), ("experience", "工作经历"),
                       ("projects", "项目经历"), ("skills", "专业技能")]:
        items = data.get(key) or []
        if not items:
            continue
        section(doc, title)
        if key == "skills":
            para(doc, " · ".join(items), size=10.5)
            continue
        for item in items:
            head = " ｜ ".join(filter(None, [item.get("company") or item.get("school") or item.get("name"),
                                             item.get("role"), item.get("degree"), item.get("period")]))
            para(doc, head, size=11, bold=True, space_after=2)
            for b in item.get("bullets", []):
                para(doc, "• " + b, size=10.5, space_after=1)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    out = guard_overwrite(args.output, args.force)
    ensure_parent(out)
    doc.save(out)
    n_exp = sum(len(i.get("bullets", [])) for i in data.get("experience", []))
    print(f"[lls-jianli] 简历已生成 → {out}（经历 {len(data.get('experience', []))} 段 / 要点 {n_exp} 条）")
    unquant = [b for i in data.get("experience", []) for b in i.get("bullets", [])
               if not any(ch.isdigit() for ch in b)]
    if unquant:
        print(f"[lls-jianli] ⚠️ {len(unquant)} 条经历无量化数字——建议与用户确认后补强，不要代为编造")


if __name__ == "__main__":
    main()
