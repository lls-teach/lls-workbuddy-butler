#!/usr/bin/env python3
"""投稿前预检：结构完整性 / 图表编号 / GB/T 7714 参考文献 / 格式残留。输出 Markdown 报告。

用法: python3 lunwen_check.py 输入.docx -o 预检报告.md
"""
import argparse
import re

from lls_lunwen_utils import die, ensure_parent

try:
    from docx import Document
except ImportError:
    die("缺少 python-docx，请先安装依赖：pip install -r scripts/requirements.txt")

REF_PAT = re.compile(r"^\[\d+\]\s+.+\[[JMCDEB/OL]+\]", re.S)
CAPTION_PAT = re.compile(r"^(图|表)\s*(\d+)-(\d+)[\s　]")
CITE_PAT = re.compile(r"(图|表)\s*(\d+-\d+)")
SENTENCE_END = ("。", "；", "？", "！")


def is_caption(text: str) -> bool:
    """题注段：以“图/表 章-序号”开头且不是完整句子（带句末标点的是正文引用句，不是题注）。"""
    return bool(CAPTION_PAT.match(text)) and not any(p in text for p in SENTENCE_END)

def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("-o", "--output", required=True)
    args = ap.parse_args()
    try:
        doc = Document(args.input)
    except Exception as e:
        die(f"无法打开: {e}")

    full = "\n".join(p.text for p in doc.paragraphs)
    problems, passed = [], []

    # 1) 结构完整性
    for kw in ("摘要", "Abstract", "关键词", "结论", "参考文献"):
        (passed if kw in full else problems).append(f"{'✅' if kw in full else '❌'} 包含「{kw}」")
    nums = [tuple(int(x) for x in m.group(1).split(".")) for p in doc.paragraphs
            if (m := re.match(r"^(\d+(?:\.\d+){0,2})[\s　]", p.text.strip()))]
    if nums:
        skipped = False
        for a, b in zip(nums, nums[1:]):
            if len(b) > len(a) + 1 or (len(b) == len(a) and b[-1] - a[-1] > 1):
                problems.append(f"❌ 章节跳级/缺章：{'.'.join(map(str, a))} → {'.'.join(map(str, b))}")
                skipped = True
        if not skipped:
            passed.append(f"✅ 章节序号 {len(nums)} 个，已检查连续性")

    # 2) 图表编号：只把题注段计为编号来源，正文引用句不计入（否则“题注+引用”会被误报重复）。
    captions = []  # (图/表, 章号, 序号)
    cites = []  # 正文引用到的编号
    for p in doc.paragraphs:
        t = p.text.strip()
        if is_caption(t):
            m = CAPTION_PAT.match(t)
            captions.append((m.group(1), int(m.group(2)), int(m.group(3))))
        else:
            for kind, num in CITE_PAT.findall(p.text):
                cites.append((kind, num))
    fig_problems = 0
    seen = {}
    for cap in captions:
        seen[cap] = seen.get(cap, 0) + 1
    for (kind, chap, num), n in sorted(seen.items()):
        if n > 1:
            problems.append(f"❌ {kind}编号重复: {chap}-{num}（出现 {n} 次题注）")
            fig_problems += 1
    for kind in ("图", "表"):
        by_chap = {}
        for k, chap, num in captions:
            if k == kind:
                by_chap.setdefault(chap, set()).add(num)
        for chap, nums_set in sorted(by_chap.items()):
            missing = [i for i in range(1, max(nums_set) + 1) if i not in nums_set]
            if missing:
                problems.append(f"❌ {kind} {chap} 章题注缺号: {missing}")
                fig_problems += 1
    caption_set = {(k, f"{c}-{n}") for k, c, n in captions}
    dangling = sorted({f"{k} {num}" for k, num in cites if (k, num) not in caption_set})
    if dangling:
        problems.append(f"❌ 正文引用无对应题注: {dangling}")
        fig_problems += 1
    if captions and fig_problems == 0:
        passed.append(f"✅ 图/表题注 {len(captions)} 个：无重复、无缺号、引用均有对应题注")
    elif not captions and not cites:
        passed.append("✅ 未发现图/表编号（无图/表时正常）")

    # 3) 参考文献 GB/T 7714
    ref_zone = full[full.rfind("参考文献"):] if "参考文献" in full else ""
    refs = [l.strip() for l in ref_zone.split("\n") if re.match(r"^\[\d+\]", l.strip())]
    bad_refs = [l[:60] for l in refs if not REF_PAT.match(l)]
    if refs:
        seq = [int(m.group(1)) for l in refs if (m := re.match(r"^\[(\d+)\]", l))]
        if seq != list(range(1, len(seq) + 1)):
            problems.append(f"❌ 参考文献序号不连续: {seq[:20]}")
        else:
            passed.append(f"✅ 参考文献序号 1-{len(seq)} 连续")
        if bad_refs:
            problems.append(f"❌ {len(bad_refs)} 条参考文献缺 GB/T 7714 要素（类型标识[J]/[M]等）: {bad_refs[:3]}")
        else:
            passed.append(f"✅ {len(refs)} 条参考文献均含类型标识")
    else:
        problems.append("❌ 未找到 [1] 式参考文献条目")

    # 4) 格式残留
    wide_spaces = sum(1 for p in doc.paragraphs if p.text.startswith("  ") and not p.text.startswith("　"))
    if wide_spaces:
        problems.append(f"❌ {wide_spaces} 个段落用空格手动缩进（应使用首行缩进样式）")
    run_empty = cur_empty = 0
    for p in doc.paragraphs:
        if not p.text.strip():
            cur_empty += 1
            run_empty = max(run_empty, cur_empty)
        else:
            cur_empty = 0
    if run_empty >= 2:
        problems.append(f"❌ 连续空行 {run_empty} 个（应删除多余空行）")

    lines = ["# 论文投稿预检报告", ""]
    lines += [f"- {p}" for p in passed] + ([f"- {p}" for p in problems] if problems else ["- ✅ 未发现问题项"])
    lines += ["", "> 本报告只覆盖程序可检查项；选题/内容/查重以学校与导师意见为准。"]
    ensure_parent(args.output)
    open(args.output, "w", encoding="utf-8").write("\n".join(lines))
    print(f"[lls-lunwen] 预检完成：通过 {len(passed)} 项，问题 {len(problems)} 项 → {args.output}")
    for p in problems:
        print("  " + p)

main()
