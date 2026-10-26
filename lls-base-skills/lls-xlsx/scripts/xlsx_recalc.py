#!/usr/bin/env python3
"""公式重算（写完公式必跑）：LibreOffice headless 就地重算并回写缓存值。

用法: python3 xlsx_recalc.py 输出.xlsx [--timeout 60]
原理: openpyxl 写公式只有字符串、没有缓存值——不重算的话 pandas/预览器/WPS 读到的公式单元格全是空。
输出: JSON 风格报告 {status, total_formulas, error_cells}。status=errors_found 时列出错误单元格。
注意: 重算通过只证明公式"算得通"，不证明"算得对"——范围错位照样零报错。
"""
import argparse
import glob
import json
import shutil
import subprocess
import sys
import tempfile
import os

from openpyxl import load_workbook

from lls_xlsx_utils import die


def soffice_or_die() -> str:
    path = shutil.which("soffice") or shutil.which("libreoffice")
    if path:
        return path
    for extra in ["/Applications/LibreOffice.app/Contents/MacOS/soffice",
                  "C:/Program Files/LibreOffice/program/soffice.exe"]:
        if os.path.exists(extra):
            return extra
    die("未找到 LibreOffice（soffice）。安装指引:\n"
        "  macOS: brew install --cask libreoffice\n"
        "  Windows/Linux: https://www.libreoffice.org/download\n"
        "安装后重跑本命令。跳过重算的后果：公式单元格在所有读取器里都是空值。")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("--timeout", type=int, default=60)
    args = ap.parse_args()

    if not os.path.exists(args.input):
        die(f"文件不存在: {args.input}")

    # 写公式统计（重算前，data_only=False）
    wb = load_workbook(args.input, data_only=False)
    formulas = sum(1 for ws in wb.worksheets for row in ws.iter_rows()
                   for c in row if isinstance(c.value, str) and c.value.startswith("="))
    wb.close()
    if formulas == 0:
        print(json.dumps({"status": "no_formulas", "total_formulas": 0}, ensure_ascii=False))
        return

    soffice = soffice_or_die()
    tmp = tempfile.mkdtemp(prefix="lls_xlsx_recalc_")
    r = subprocess.run([soffice, "--headless", "--convert-to", "xlsx", "--outdir", tmp, args.input],
                       capture_output=True, text=True, timeout=args.timeout)
    converted = glob.glob(os.path.join(tmp, "*.xlsx"))
    if not converted:
        die(f"soffice 重算失败: {r.stderr[:300]}")
    shutil.move(converted[0], args.input)
    shutil.rmtree(tmp, ignore_errors=True)

    # 读回检查错误值
    wb = load_workbook(args.input, data_only=True)
    error_cells = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("#") and c.value.endswith(("!", "?", "A", "0")):
                    error_cells.append(f"{ws.title}!{c.coordinate}={c.value}")
    wb.close()

    status = "errors_found" if error_cells else "success"
    report = {"status": status, "total_formulas": formulas, "error_count": len(error_cells)}
    if error_cells:
        report["error_cells"] = error_cells[:100]
    print(json.dumps(report, ensure_ascii=False))
    if status == "errors_found":
        print("[lls-xlsx] ⚠️ 存在公式错误单元格，逐个修复后重跑本命令", file=sys.stderr)
        sys.exit(0)  # errors_found 不算脚本失败：修复责任在公式本身
    print(f"[lls-xlsx] ✅ {formulas} 个公式已重算并回写缓存值 → {args.input}")


if __name__ == "__main__":
    main()
