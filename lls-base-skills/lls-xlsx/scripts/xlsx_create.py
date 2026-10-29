#!/usr/bin/env python3
"""创建工作簿：从 CSV/JSON 装数据，支持表头样式、SUM 公式列、图表、中文列宽自适应。

用法:
  python3 xlsx_create.py -o 报表.xlsx --from-csv 数据.csv --title 10月报销 --sum-col 金额 --chart bar
  python3 xlsx_create.py -o 新表.xlsx --from-json rows.json
  手动模式（空表+表头）: python3 xlsx_create.py -o 新表.xlsx --headers 姓名,部门,金额
"""
import argparse
import json
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from lls_xlsx_utils import cjk_width, die, ensure_parent, guard_overwrite, smart_read_csv


def autofit(ws) -> None:
    for col_cells in ws.columns:
        width = max((cjk_width(c.value) for c in col_cells if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(width + 2, 8), 60)


def add_chart(ws, kind: str, ncols: int, nrows: int) -> None:
    if ncols < 2 or nrows < 2:
        print("[lls-xlsx] ⚠️ 数据不足两行两列，跳过图表", file=sys.stderr)
        return
    cls = {"bar": BarChart, "line": LineChart, "pie": PieChart}.get(kind)
    if not cls:
        return
    chart = cls()
    chart.title = ws.title
    data = Reference(ws, min_col=ncols, min_row=1, max_row=nrows)
    cats = Reference(ws, min_col=1, min_row=2, max_row=nrows)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.width, chart.height = 16, 9
    ws.add_chart(chart, f"{get_column_letter(ncols + 2)}2")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--output", required=True)
    ap.add_argument("--from-csv")
    ap.add_argument("--from-json", help='JSON 数组，元素为对象或数组')
    ap.add_argument("--headers", help="手动模式表头: 姓名,部门,金额")
    ap.add_argument("--title", default="Sheet1")
    ap.add_argument("--sum-col", help="对该列名末行追加 SUM 公式")
    ap.add_argument("--chart", choices=["bar", "line", "pie"])
    ap.add_argument("--style", action="store_true", default=True, help="表头加样式（默认开）")
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()

    if args.from_csv:
        matrix = smart_read_csv(args.from_csv)
    elif args.from_json:
        with open(args.from_json, encoding="utf-8") as f:
            data = json.load(f)
        if data and isinstance(data[0], dict):
            headers = list(data[0].keys())
            matrix = [headers] + [[str(r.get(h, "")) for h in headers] for r in data]
        else:
            matrix = [[str(c) for c in r] for r in data]
    elif args.headers:
        matrix = [[h.strip() for h in args.headers.split(",")]]
    else:
        die("需要 --from-csv / --from-json / --headers 之一")

    if not matrix:
        die("输入为空：CSV/JSON 无数据行，请检查输入文件是否为空")

    def coerce(v: str):
        """数值型字符串转数字（否则 SUM/COUNT 类公式按 0 处理）。"""
        if not isinstance(v, str) or not v.strip():
            return v
        s = v.strip().replace(",", "")
        try:
            return int(s)
        except ValueError:
            pass
        try:
            return float(s)
        except ValueError:
            return v

    if args.from_csv or args.from_json:
        matrix = [matrix[0]] + [[coerce(c) for c in row] for row in matrix[1:]]

    wb = Workbook()
    ws = wb.active
    ws.title = args.title or "Sheet1"
    for row in matrix:
        ws.append(row)

    if args.style and matrix:
        fill = PatternFill("solid", fgColor="1F6FEB")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center")

    if args.sum_col and matrix and len(matrix) > 1:
        headers = matrix[0]
        if args.sum_col in headers:
            col_idx = headers.index(args.sum_col) + 1
            letter = get_column_letter(col_idx)
            n = len(matrix)
            ws.cell(row=n + 1, column=col_idx, value=f"=SUM({letter}2:{letter}{n})")
            ws.cell(row=n + 1, column=1, value="合计")
            print(f"[lls-xlsx] 已加合计行: {letter}{n + 1} =SUM({letter}2:{letter}{n})")
        else:
            print(f"[lls-xlsx] ⚠️ --sum-col {args.sum_col!r} 不在表头 {headers} 中", file=sys.stderr)

    autofit(ws)
    if args.chart:
        add_chart(ws, args.chart, len(matrix[0]), len(matrix))

    out = guard_overwrite(args.output, args.force)
    ensure_parent(out)
    wb.save(out)
    print(f"[lls-xlsx] 已创建 → {out}（{len(matrix)} 行 x {len(matrix[0])} 列）")


if __name__ == "__main__":
    main()
