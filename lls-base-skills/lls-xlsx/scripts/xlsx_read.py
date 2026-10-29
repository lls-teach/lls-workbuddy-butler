#!/usr/bin/env python3
"""读取 xlsx：打印结构概览，可选输出 JSON/CSV。

用法: python3 xlsx_read.py 输入.xlsx [--sheet Sheet1] [--json out.json] [--to-csv out_dir/] [--rows 10]
"""
import argparse
import csv
import json
import os

from openpyxl import load_workbook

from lls_xlsx_utils import die, ensure_parent


def sheet_matrix(ws, max_rows: int | None = None) -> list[list]:
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if max_rows and i >= max_rows:
            break
        rows.append(["" if v is None else v for v in row])
    return rows


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("--sheet", help="工作表名（缺省=第一个）")
    ap.add_argument("--json", help="把数据写成 JSON 文件")
    ap.add_argument("--to-csv", help="把数据写成 CSV 文件")
    ap.add_argument("--rows", type=int, default=8, help="stdout 预览行数（默认 8，0=只报结构）")
    args = ap.parse_args()

    try:
        wb = load_workbook(args.input, data_only=True, read_only=True)
    except Exception as e:  # noqa: BLE001
        die(f"无法打开（若是 .xls 旧格式请先另存为 .xlsx）: {e}")

    names = wb.sheetnames
    target = args.sheet or names[0]
    if target not in names:
        die(f"工作表 {target!r} 不存在，可选: {names}")
    ws = wb[target]
    matrix = sheet_matrix(ws, None if args.json or args.to_csv else (args.rows or None))

    print(f"[lls-xlsx] 文件: {args.input}")
    print(f"[lls-xlsx] 工作表: {names}（当前: {target}，{ws.max_row} 行 x {ws.max_column} 列）")
    if args.rows:
        for row in matrix[: args.rows]:
            print("  " + " | ".join(str(v)[:18] for v in row))

    if args.json:
        ensure_parent(args.json)
        header = matrix[0] if matrix else []
        data = [dict(zip(header, r)) for r in matrix[1:]] if header else matrix
        with open(args.json, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=1, default=str)
        print(f"[lls-xlsx] JSON → {args.json}（{len(data)} 条记录）")
    if args.to_csv:
        ensure_parent(os.path.join(args.to_csv, "_"))
        out = os.path.join(args.to_csv, f"{target}.csv")
        with open(out, "w", newline="", encoding="utf-8-sig") as f:
            csv.writer(f).writerows(matrix)
        print(f"[lls-xlsx] CSV → {out}")
    wb.close()


if __name__ == "__main__":
    main()
