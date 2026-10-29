#!/usr/bin/env python3
"""创建中文 Word 文档（标题层级/列表/引用/表格/页眉）。

用法:
  printf '标题\n\n正文第一段\n\n- 要点1\n- 要点2\n\n[TABLE]\n姓名\t部门\n张三\t市场部\n\n落款：行政部\n' | \
    python3 docx_create.py -o 输出.docx --title 文档标题 --header 页眉文字
"""
import argparse
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from lls_docx_utils import CN_FONT_HEI, die, ensure_parent, guard_overwrite, set_cn_font


def add_para(doc, text: str, size: int = 12, bold: bool = False, align=None, style=None) -> None:
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    set_cn_font(run)
    if align is not None:
        p.alignment = align


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--output", required=True)
    ap.add_argument("--title", default="")
    ap.add_argument("--header", default="")
    ap.add_argument("--stdin-file")
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()

    raw = None
    if args.stdin_file:
        raw = open(args.stdin_file, encoding="utf-8").read()
    else:
        # Windows 控制台默认编码未必是 UTF-8，管道输入中文会乱码；显式切到 UTF-8
        try:
            sys.stdin.reconfigure(encoding="utf-8")
        except Exception:  # noqa: BLE001 - 旧解释器/非常驻 stdin 时保持原编码
            pass
        raw = sys.stdin.read()
    if not raw.strip():
        die("正文为空：从 stdin 或 --stdin-file 提供内容")

    doc = Document()
    if args.header:
        hdr = doc.sections[0].header.paragraphs[0]
        run = hdr.add_run(args.header)
        run.font.size = Pt(9)
        set_cn_font(run)
        hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if args.title:
        add_para(doc, args.title, size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    lines = raw.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("[TABLE]"):
            rows = []
            i += 1
            while i < len(lines) and lines[i].strip():
                rows.append(lines[i].split("\t"))
                i += 1
            if rows:
                # 列数取各行最大值：后行比首行多列时不越界，短行留空单元格
                ncols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=ncols)
                table.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c, cell in enumerate(row):
                        cell_p = table.cell(r, c).paragraphs[0]
                        run = cell_p.add_run(cell.strip())
                        run.font.size = Pt(10.5)
                        set_cn_font(run)
                        if r == 0:
                            run.font.bold = True
            continue
        if line.startswith("# "):
            add_para(doc, line[2:], size=15, bold=True)
        elif line.startswith("## "):
            add_para(doc, line[3:], size=13, bold=True)
        elif line.startswith("- "):
            add_para(doc, line[2:], style="List Bullet")
        elif line.startswith("> "):
            add_para(doc, line[2:], size=10.5)
        else:
            add_para(doc, line)
        i += 1

    out = guard_overwrite(args.output, args.force)
    ensure_parent(out)
    doc.save(out)
    print(f"[lls-docx] 已创建 → {out}")


if __name__ == "__main__":
    main()
