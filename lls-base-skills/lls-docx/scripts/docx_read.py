#!/usr/bin/env python3
"""读取 docx 全文（段落+表格），可选 JSON。

用法: python3 docx_read.py 输入.docx [--json out.json]
"""
import argparse
import json

from docx import Document

from lls_docx_utils import die, ensure_parent


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("--json")
    args = ap.parse_args()

    try:
        doc = Document(args.input)
    except Exception as e:  # noqa: BLE001
        die(f"无法打开（.doc 旧格式请先另存为 .docx）: {e}")

    blocks = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            blocks.append({"type": "para", "style": para.style.name, "text": t})
    for t_no, table in enumerate(doc.tables, 1):
        rows = [[c.text.strip() for c in row.cells] for row in table.rows]
        blocks.append({"type": "table", "index": t_no, "rows": rows})
        print(f"===== 表格 {t_no}（{len(rows)} 行）=====")
        for row in rows[:3]:
            print("  " + " | ".join(row))

    for b in blocks:
        if b["type"] == "para":
            print(b["text"][:100])

    print(f"[lls-docx] 段落+表格共 {len(blocks)} 块")
    if args.json:
        ensure_parent(args.json)
        json.dump(blocks, open(args.json, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        print(f"[lls-docx] JSON → {args.json}")


if __name__ == "__main__":
    main()
