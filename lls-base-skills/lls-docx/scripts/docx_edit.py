#!/usr/bin/env python3
"""编辑 docx：批量替换文字 / 追加段落。默认不动原文件。

用法:
  替换: python3 docx_edit.py 输入.docx --replace "甲方=委托方" -o 新.docx
  追加: python3 docx_edit.py 输入.docx --append "补充条款：本协议自签订之日起生效。" -o 新.docx
"""
import argparse

from docx import Document

from lls_docx_utils import die, ensure_parent, guard_overwrite, set_cn_font


def replace_in_paragraph(para, pairs: dict[str, str]) -> int:
    count = 0
    for run in para.runs:
        for old, new in pairs.items():
            if old in run.text:
                count += run.text.count(old)
                run.text = run.text.replace(old, new)
    return count


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("--replace", dest="replacements", action="append", default=[], help="旧=新，可多次")
    ap.add_argument("--append", help="在文末追加一个段落")
    ap.add_argument("-o", "--output", required=True)
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()

    try:
        doc = Document(args.input)
    except Exception as e:  # noqa: BLE001
        die(f"无法打开: {e}")

    total = 0
    if args.replacements:
        pairs = dict(r.split("=", 1) for r in args.replacements)
        for para in doc.paragraphs:
            total += replace_in_paragraph(para, pairs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        total += replace_in_paragraph(para, pairs)
        print(f"[lls-docx] 替换完成，共 {total} 处")
        if total == 0:
            print("[lls-docx] ⚠️ 0 处命中——请核对替换词；文本框/页眉中的文字不在此能力范围")

    if args.append:
        para = doc.add_paragraph()
        run = para.add_run(args.append)
        set_cn_font(run)
        print(f"[lls-docx] 已追加段落: {args.append[:40]}…")

    if not args.replacements and not args.append:
        die("至少给一个 --replace 或 --append")

    out = guard_overwrite(args.output, args.force)
    ensure_parent(out)
    doc.save(out)
    print(f"[lls-docx] 已保存 → {out}（原文件未动）")


if __name__ == "__main__":
    main()
