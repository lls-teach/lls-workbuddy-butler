"""lls-shujv 共享工具：零 pandas 的表统计（csv/openpyxl 双读 + 数值列识别 + IQR 异常）。"""
import csv
import os
import statistics
import sys


def die(msg, code=1):
    print(f"[lls-shujv] 错误: {msg}", file=sys.stderr)
    sys.exit(code)


def ensure_parent(path):
    parent = os.path.dirname(os.path.abspath(path))
    os.makedirs(parent, exist_ok=True)
    return path


def read_table(path: str) -> tuple[list[str], list[list]]:
    """读 CSV/Excel → (表头, 行数据)；数值字符串自动转数字。"""
    if path.lower().endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ModuleNotFoundError:
            die("缺少 openpyxl：请先按 SKILL.md 第 0 步安装 scripts/requirements.txt")
        try:
            wb = load_workbook(path, data_only=True, read_only=True)
        except OSError as e:
            die(f"无法读取文件: {path} ({e})")
        ws = wb[wb.sheetnames[0]]
        rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
        wb.close()
    elif path.lower().endswith(".csv"):
        for enc in ("utf-8-sig", "utf-8", "gbk"):
            try:
                with open(path, encoding=enc, newline="") as f:
                    rows = [list(r) for r in csv.reader(f)]
                break
            except UnicodeDecodeError:
                continue
            except OSError as e:
                die(f"无法读取文件: {path} ({e})")
        else:
            die(f"CSV 编码无法识别: {path}")
    else:
        die("仅支持 .csv / .xlsx")
    if not rows:
        die("空表")
    header = [str(h).strip() for h in rows[0]]
    data = []
    for r in rows[1:]:
        row = []
        for v in r:
            v = str(v).strip() if v is not None else ""
            try:
                row.append(float(v.replace(",", "")))
            except ValueError:
                row.append(v)
        data.append(row)
    return header, data


def col_index(header: list[str], name: str) -> int:
    for i, h in enumerate(header):
        if h == name:
            return i
    die(f"列 {name!r} 不存在，可选: {header}")


def numeric_cols(header, data) -> list[int]:
    out = []
    for i in range(len(header)):
        vals = [r[i] for r in data if i < len(r)]
        nums = [v for v in vals if isinstance(v, (int, float))]
        if nums and len(nums) / max(len(vals), 1) > 0.8:
            out.append(i)
    return out


def describe(vals: list[float]) -> dict:
    vals = sorted(vals)
    n = len(vals)
    q1 = vals[int(n * 0.25)] if n else 0
    q3 = vals[int(n * 0.75)] if n else 0
    iqr = q3 - q1
    outliers = [v for v in vals if v < q1 - 1.5 * iqr or v > q3 + 1.5 * iqr] if iqr else []
    return {"count": n, "sum": round(sum(vals), 2), "mean": round(statistics.fmean(vals), 2) if n else 0,
            "min": vals[0] if n else 0, "max": vals[-1] if n else 0,
            "outliers": outliers[:10], "outlier_n": len(outliers)}


def group_agg(data, key_idx: int, val_idx: int, agg: str = "sum") -> list[tuple]:
    groups: dict[str, list[float]] = {}
    for r in data:
        k = str(r[key_idx]) if key_idx < len(r) else "?"
        v = r[val_idx] if val_idx < len(r) else None
        if isinstance(v, (int, float)):
            groups.setdefault(k, []).append(float(v))
    fn = {"sum": sum, "mean": lambda vs: round(statistics.fmean(vs), 2), "count": len}.get(agg)
    if agg == "count":
        return sorted(((k, len(vs)) for k, vs in groups.items()), key=lambda kv: -kv[1])
    return sorted(((k, round(fn(vs), 2)) for k, vs in groups.items()), key=lambda kv: -kv[1])


def guard_overwrite(path: str, force: bool = False) -> str:
    if not os.path.exists(path) or force:
        return path
    root, ext = os.path.splitext(path)
    alt = f"{root}_new{ext}"
    print(f"[lls-shujv] 目标已存在，改写到 {alt}", file=sys.stderr)
    return alt
