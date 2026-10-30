"""构造 pptx_check 的正负例，验证封面豁免没有把内容页的检查一并放过。"""
import pathlib
import sys

from pptx import Presentation
from pptx.util import Inches, Pt

OUT = pathlib.Path(r"F:\CODE\skill_windows\real-test\runs\lls-pptx\检查器用例")
OUT.mkdir(parents=True, exist_ok=True)


def add_slide(prs, title, bullets, title_size=28, body_size=16):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(9), Inches(1.0))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = title
    run.font.size = Pt(title_size)
    if bullets:
        body = slide.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(9), Inches(4))
        tf = body.text_frame
        tf.word_wrap = True
        for idx, item in enumerate(bullets):
            para = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            r = para.add_run()
            r.text = "• " + item
            r.font.size = Pt(body_size)
    return slide


# 用例 A：正常演示——封面 + 内容页（含应报与不应报）
prs = Presentation()
add_slide(prs, "数据看板入门", ["2026-10-30", "运营组内部分享"])
add_slide(prs, "市场分析", ["市场规模", "竞争格局"])                      # 应报：话题词
add_slide(prs, "1.", ["第一点", "第二点"])                                # 应报：纯编号
add_slide(prs, "背景介绍", ["起因", "经过"])                              # 应报：话题词
add_slide(prs, "三日营收增长 9.9%，增量来自客户数", ["企业客户 +9.2%", "中小企业 +10.9%"])  # 不应报
add_slide(prs, "三个关键动作", ["定口径", "存快照", "选图表"])            # 应报：内容页短话题标题
p = OUT / "用例A-正常演示.pptx"
prs.save(p)
print("已生成", p)

# 用例 B：开门就上内容（第 1 页满版），短标题不应被当封面豁免
prs = Presentation()
add_slide(
    prs,
    "市场分析",
    [
        "企业客户日营收三天从 12.80 万升至 13.98 万，累计增长 9.2%，未见单日回撤",
        "中小企业日营收三天从 8.60 万升至 9.54 万，累计增长 10.9%，增速高于企业客户",
        "增量几乎全部来自成交客户数增加，客单价小幅回落，增长结构偏多签单而非签大单",
        "两个客群合计日营收从 21.40 万升至 23.52 万，三日累计增长 9.9%，与目标持平",
    ],
)
p = OUT / "用例B-首屏满版.pptx"
prs.save(p)
print("已生成", p)
