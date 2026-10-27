#!/usr/bin/env python3
"""Generate all test fixtures into tests/fixtures/. Idempotent."""
import os, sys, subprocess, shutil, wave, math, struct, zipfile, tarfile

FIX = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures")
os.makedirs(FIX, exist_ok=True)
TOKEN = "LHGCONVERT"
CN = "格式转换测试"
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "scripts"))

def w(name, content, mode="w"):
    p = os.path.join(FIX, name)
    with open(p, mode, encoding=None if "b" in mode else "utf-8") as fh: fh.write(content)
    return p

def soffice(src, tgt, outdir=FIX):
    r = subprocess.run(["soffice","--headless","--norestore","--convert-to",tgt,"--outdir",outdir,src], capture_output=True, text=True, timeout=180)
    return r

def ffmpeg(args):
    r = subprocess.run(["ffmpeg","-y","-v","error"]+args, capture_output=True, text=True, timeout=180)
    if r.returncode != 0: raise RuntimeError("ffmpeg fixture failed: "+r.stderr[-200:])

# --- text family ---
txt_body = f"{TOKEN} {CN}\n第二行内容，第二段。\n数字 12345。\n"
w("sample.txt", txt_body)
w("sample.md", f"# {TOKEN} 标题\n\n{CN} 正文段落，包含 **加粗** 与列表：\n\n- 项目一\n- 项目二\n\n结尾段落。\n")
w("sample.html", f"<!DOCTYPE html><html><head><meta charset='utf-8'><title>{TOKEN}</title></head><body><h1>{TOKEN} {CN}</h1><p>正文段落 abc 12345。</p><ul><li>项目一</li><li>项目二</li></ul></body></html>")

# --- docx / xlsx / pptx via python libs ---
from docx import Document
d = Document(); d.add_heading(f"{TOKEN} {CN}", level=1)
d.add_paragraph(f"正文段落包含标记 {TOKEN} 与中文{CN}，数字 12345。")
d.add_paragraph("第二段内容。")
d.save(os.path.join(FIX, "sample.docx"))

from openpyxl import Workbook
wb = Workbook(); ws = wb.active
ws.append(["name", "value", "note"])
ws.append([TOKEN, 12345, CN])
ws.append(["第二行", 678, "数据"])
wb.save(os.path.join(FIX, "sample.xlsx"))
w("sample.csv", f"name,value,note\n{TOKEN},12345,{CN}\n第二行,678,数据\n")
w("sample.tsv", f"name\tvalue\tnote\n{TOKEN}\t12345\t{CN}\n第二行\t678\t数据\n")
import json, yaml
recs = [{"name": TOKEN, "value": 12345, "note": CN}, {"name": "第二行", "value": 678, "note": "数据"}]
w("sample.json", json.dumps(recs, ensure_ascii=False, indent=2))
w("sample.yaml", yaml.safe_dump(recs, allow_unicode=True))
rows_xml = "".join(f"<row><name>{r['name']}</name><value>{r['value']}</value><note>{r['note']}</note></row>" for r in recs)
w("sample.xml", f"<?xml version='1.0' encoding='utf-8'?><data>{rows_xml}</data>")

from pptx import Presentation
prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = f"{TOKEN} {CN}"
slide.placeholders[1].text = "要点一：格式转换\n要点二：本地运行"
prs.save(os.path.join(FIX, "sample.pptx"))

# --- images ---
from PIL import Image, ImageDraw
im = Image.new("RGB", (320, 200), (30, 90, 160))
dr = ImageDraw.Draw(im); dr.rectangle([20, 20, 300, 180], outline=(255, 255, 0), width=4); dr.text((40, 90), TOKEN, fill=(255, 255, 255))
im.save(os.path.join(FIX, "sample.png")); im.save(os.path.join(FIX, "sample.jpg"))
im.save(os.path.join(FIX, "sample.webp"))
w("sample.svg", f'<svg xmlns="http://www.w3.org/2000/svg" width="320" height="200"><rect width="320" height="200" fill="#305aa0"/><text x="40" y="100" fill="#ffffff" font-size="20">{TOKEN}</text></svg>')

# --- audio / video via ffmpeg ---
wav_path = os.path.join(FIX, "sample.wav")
with wave.open(wav_path, "wb") as wf:
    wf.setnchannels(1); wf.setsampwidth(2); wf.setframerate(22050)
    frames = b"".join(struct.pack("<h", int(8000*math.sin(2*math.pi*440*i/22050))) for i in range(22050))
    wf.writeframes(frames)
for ext in ("mp3", "flac", "ogg", "m4a"):
    ffmpeg(["-i", wav_path, os.path.join(FIX, f"sample.{ext}")])
mp4_path = os.path.join(FIX, "sample.mp4")
ffmpeg(["-f","lavfi","-i","testsrc=duration=2:size=320x240:rate=10","-f","lavfi","-i","sine=frequency=440:duration=2","-shortest","-pix_fmt","yuv420p", mp4_path])
for ext in ("mov", "webm", "avi", "mkv"):
    ffmpeg(["-i", mp4_path, os.path.join(FIX, f"sample.{ext}")])

# --- subtitles ---
w("sample.srt", f"1\n00:00:00,000 --> 00:00:01,500\n{CN} {TOKEN}\n\n2\n00:00:01,500 --> 00:00:03,000\n第二句字幕\n")
w("sample.vtt", f"WEBVTT\n\n00:00:00.000 --> 00:00:01.500\n{CN} {TOKEN}\n\n00:00:01.500 --> 00:00:03.000\n第二句字幕\n")
sys.path.insert(0, os.path.join(FIX, "..", "..", "scripts"))
import importlib.util
spec = importlib.util.spec_from_file_location("conv", os.path.join(FIX, "..", "..", "scripts", "convert.py"))
conv = importlib.util.module_from_spec(spec); spec.loader.exec_module(conv)
cues = conv.parse_srt(open(os.path.join(FIX,"sample.srt"), encoding="utf-8").read())
w("sample.ass", conv.cues_to_ass(cues))

# --- epub via engine builder ---
paras = [f"{TOKEN} {CN}", "第一章正文段落，数字 12345。", "第二段内容结束。"]
conv.make_epub(f"{TOKEN} {CN}", paras, os.path.join(FIX, "sample.epub"))

# --- archives ---
inner = w("inner_note.txt", f"{TOKEN} {CN} 压缩包内文件\n")
zp = os.path.join(FIX, "sample.zip")
with zipfile.ZipFile(zp, "w") as z: z.write(inner, "inner_note.txt")
os.remove(inner)
tgz = os.path.join(FIX, "sample.tar.gz")
with tarfile.open(tgz, "w:gz") as t:
    p2 = w("inner_note.txt", f"{TOKEN} {CN} 压缩包内文件\n"); t.add(p2, "inner_note.txt"); os.remove(p2)
import py7zr
sz = os.path.join(FIX, "sample.7z")
p3 = w("inner_note.txt", f"{TOKEN} {CN} 压缩包内文件\n")
with py7zr.SevenZipFile(sz, "w") as z: z.writeall(os.path.join(FIX, "inner_note.txt"), "inner_note.txt")
os.remove(p3)

# --- fonts ---
ttf_src = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
shutil.copy(ttf_src, os.path.join(FIX, "sample.ttf"))
from fontTools.ttLib import TTFont
f = TTFont(ttf_src); f.save(os.path.join(FIX, "sample.otf"))
f2 = TTFont(ttf_src); f2.flavor = "woff"; f2.save(os.path.join(FIX, "sample.woff"))
f3 = TTFont(ttf_src); f3.flavor = "woff2"; f3.save(os.path.join(FIX, "sample.woff2"))

# --- libreoffice-derived fixtures ---
docx = os.path.join(FIX, "sample.docx"); xlsx = os.path.join(FIX, "sample.xlsx"); pptx = os.path.join(FIX, "sample.pptx")
for tgt in ("odt", "rtf", "doc", "pdf"):
    soffice(docx, tgt)
for tgt in ("ods", "xls"):
    soffice(xlsx, tgt)
for tgt in ("odp", "ppt"):
    soffice(pptx, tgt)

# dummy unsupported source
w("sample.pages", "dummy pages file\n")
print("fixtures:", len(os.listdir(FIX)))
print(sorted(os.listdir(FIX)))
