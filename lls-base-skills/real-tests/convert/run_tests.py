#!/usr/bin/env python3
"""lls-convert test suite: every case runs a real conversion via the engine
and validates the artifact by reading it back. Results stream to
results.jsonl; summary to last_run.txt."""
import os, sys, json, subprocess, shutil, zipfile, tarfile, tempfile, io

ROOT = os.path.dirname(os.path.abspath(__file__))   # real-tests/convert
FIX = os.path.join(ROOT, "fixtures")
OUT = os.path.join(ROOT, "out")
ENGINE = os.path.join(os.path.dirname(os.path.dirname(ROOT)), "lls-convert", "scripts", "convert.py")
TOKEN = "LHGCONVERT"
sys.path.insert(0, os.path.dirname(ENGINE))

def read_txt(p):
    with open(p, encoding="utf-8", errors="replace") as fh: return fh.read()

def v_token_text(p): assert TOKEN in read_txt(p), "token missing"
def v_pdf_text(p):
    from pypdf import PdfReader
    r = PdfReader(p); assert len(r.pages) >= 1
    assert TOKEN in "".join((pg.extract_text() or "") for pg in r.pages), "token missing in pdf text"
def v_pdf_pages(p):
    from pypdf import PdfReader
    assert len(PdfReader(p).pages) >= 1
def v_docx(p):
    from docx import Document
    d = Document(p)
    text = "\n".join(par.text for par in d.paragraphs)
    assert TOKEN in text, "token missing in docx"
def v_xlsx(p):
    from openpyxl import load_workbook
    wb = load_workbook(p)
    vals = [str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None]
    assert any(TOKEN in v for v in vals), "token missing in xlsx"
def v_pptx(p):
    from pptx import Presentation
    prs = Presentation(p)
    texts = []
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame: texts.append(sh.text_frame.text)
    assert any(TOKEN in t for t in texts), "token missing in pptx"
def v_epub(p):
    assert zipfile.is_zipfile(p)
    import importlib.util
    spec = importlib.util.spec_from_file_location("conv", ENGINE); conv = importlib.util.module_from_spec(spec); spec.loader.exec_module(conv)
    assert any(TOKEN in t for t in conv.epub_to_paragraphs(p)), "token missing in epub"
def v_odf(p):
    assert zipfile.is_zipfile(p), "odf not a zip"
    with zipfile.ZipFile(p) as z:
        content = z.read("content.xml").decode("utf-8", errors="replace")
    assert TOKEN in content, "token missing in odf content.xml"
def v_bytes(p):
    data = open(p, "rb").read()
    assert TOKEN.encode() in data or TOKEN.encode("utf-16le") in data, "token missing (bytes)"
def v_image(fmt):
    def check(p):
        from PIL import Image
        im = Image.open(p); im.verify()
        im2 = Image.open(p)
        assert im2.format == fmt, f"format {im2.format} != {fmt}"
    return check
def v_av(kind):
    def check(p):
        r = subprocess.run(["ffprobe","-v","error","-show_entries","format=duration:stream=codec_type","-of","json",p], capture_output=True, text=True, timeout=60)
        info = json.loads(r.stdout)
        dur = float(info.get("format", {}).get("duration", 0))
        assert dur > 0.3, f"duration {dur}"
        kinds = [s.get("codec_type") for s in info.get("streams", [])]
        assert kind in kinds, f"no {kind} stream in {kinds}"
    return check
def v_archive(p):
    if p.endswith(".zip"):
        with zipfile.ZipFile(p) as z: names = z.namelist()
    elif p.endswith(".tar.gz"):
        with tarfile.open(p) as t: names = t.getnames()
    elif p.endswith(".7z"):
        import py7zr
        with py7zr.SevenZipFile(p) as z: names = z.getnames()
    assert any("inner_note.txt" in n for n in names), f"inner file missing: {names}"
def v_font(p):
    from fontTools.ttLib import TTFont
    f = TTFont(p); assert f["name"] is not None
def v_sub(marker):
    def check(p):
        t = read_txt(p)
        assert "格式转换测试" in t and marker in t, f"subtitle marker {marker} missing"
    return check

CASES = []
def C(fixture, tgt, validator): CASES.append((fixture, tgt, validator))

# documents
for t, v in [("pdf", v_pdf_text), ("txt", v_token_text), ("md", v_token_text), ("html", v_token_text), ("odt", v_odf), ("rtf", v_bytes)]: C("sample.docx", t, v)
for t, v in [("docx", v_docx), ("pdf", v_pdf_text)]: C("sample.odt", t, v)
for t, v in [("docx", v_docx), ("pdf", v_pdf_pages)]: C("sample.rtf", t, v)
for t, v in [("docx", v_docx), ("pdf", v_pdf_pages)]: C("sample.doc", t, v)
for t, v in [("md", v_token_text), ("html", v_token_text), ("pdf", v_pdf_pages), ("docx", v_docx)]: C("sample.txt", t, v)
for t, v in [("txt", v_token_text), ("html", v_token_text), ("pdf", v_pdf_pages), ("docx", v_docx), ("odt", v_odf), ("epub", v_epub)]: C("sample.md", t, v)
for t, v in [("md", v_token_text), ("txt", v_token_text), ("pdf", v_pdf_pages), ("docx", v_docx)]: C("sample.html", t, v)
for t, v in [("txt", v_token_text), ("md", v_token_text), ("docx", v_docx), ("png", v_image("PNG")), ("jpg", v_image("JPEG"))]: C("sample.pdf", t, v)
for t, v in [("txt", v_token_text), ("md", v_token_text), ("docx", v_docx), ("pdf", v_pdf_pages)]: C("sample.epub", t, v)
# data
for t in ("csv","tsv","json","xml","yaml","md","html"): C("sample.xlsx", t, v_token_text)
C("sample.xlsx", "ods", v_odf); C("sample.xlsx", "pdf", v_pdf_pages)
for t, v in [("xlsx", v_xlsx), ("json", v_token_text), ("xml", v_token_text), ("yaml", v_token_text), ("md", v_token_text), ("pdf", v_pdf_pages)]: C("sample.csv", t, v)
for t, v in [("csv", v_token_text), ("xlsx", v_xlsx), ("xml", v_token_text), ("yaml", v_token_text)]: C("sample.json", t, v)
for t, v in [("csv", v_token_text), ("json", v_token_text)]: C("sample.xml", t, v)
for t, v in [("csv", v_token_text), ("json", v_token_text), ("xlsx", v_xlsx)]: C("sample.yaml", t, v)
for t, v in [("csv", v_token_text), ("json", v_token_text)]: C("sample.tsv", t, v)
for t, v in [("xlsx", v_xlsx), ("csv", v_token_text), ("pdf", v_pdf_pages)]: C("sample.ods", t, v)
for t, v in [("xlsx", v_xlsx), ("csv", v_token_text)]: C("sample.xls", t, v)
# slides
for t, v in [("pdf", v_pdf_text), ("odp", v_odf), ("png", v_image("PNG")), ("jpg", v_image("JPEG"))]: C("sample.pptx", t, v)
for t, v in [("pptx", v_pptx), ("pdf", v_pdf_pages)]: C("sample.ppt", t, v)
for t, v in [("pptx", v_pptx), ("pdf", v_pdf_pages)]: C("sample.odp", t, v)
# images
for t, fmt in [("jpg","JPEG"), ("webp","WEBP"), ("bmp","BMP"), ("tiff","TIFF"), ("ico","ICO"), ("gif","GIF")]: C("sample.png", t, v_image(fmt))
C("sample.png", "pdf", v_pdf_pages)
for t, v in [("png", v_image("PNG")), ("webp", v_image("WEBP")), ("pdf", v_pdf_pages)]: C("sample.jpg", t, v)
C("sample.webp", "png", v_image("PNG"))
for t, v in [("png", v_image("PNG")), ("pdf", v_pdf_pages)]: C("sample.svg", t, v)
# audio / video
for t in ("mp3","flac","ogg","m4a","aac"): C("sample.wav", t, v_av("audio"))
for t in ("wav","ogg","flac"): C("sample.mp3", t, v_av("audio"))
for t in ("webm","mov","avi","mkv"): C("sample.mp4", t, v_av("video"))
for t in ("mp3","wav"): C("sample.mp4", t, v_av("audio"))
C("sample.mp4", "gif", v_image("GIF")); C("sample.mp4", "png", v_image("PNG"))
for f in ("avi","mkv","webm","mov"): C(f"sample.{f}", "mp4", v_av("video"))
# subtitles
C("sample.srt", "vtt", v_sub("WEBVTT")); C("sample.srt", "ass", v_sub("Dialogue:"))
C("sample.vtt", "srt", v_sub("-->")); C("sample.vtt", "ass", v_sub("Dialogue:"))
C("sample.ass", "srt", v_sub("-->")); C("sample.ass", "vtt", v_sub("WEBVTT"))
# archives
C("sample.zip", "tar.gz", v_archive); C("sample.zip", "7z", v_archive)
C("sample.tar.gz", "zip", v_archive); C("sample.7z", "zip", v_archive)
# fonts
C("sample.ttf", "woff", v_font); C("sample.ttf", "woff2", v_font)
C("sample.otf", "woff2", v_font); C("sample.otf", "woff", v_font)
C("sample.woff", "ttf", v_font); C("sample.woff2", "ttf", v_font)

NEGATIVES = [("sample.ttf","mp4"), ("sample.mp3","docx"), ("sample.pages","docx"),
             ("sample.pdf","xlsx"), ("sample.svg","mp3"), ("sample.zip","pdf"), ("sample.docx","mp4")]

# ---------- 安全性用例 ----------
# 这三类不是「格式对不对」，而是「会不会毁用户文件」，单独跑、单独断言。
def _mk_tar_with_member(path, member, payload):
    import io
    with tarfile.open(path, "w:gz") as t:
        ti = tarfile.TarInfo(member); ti.size = len(payload)
        t.addfile(ti, io.BytesIO(payload))

def _mk_tar_with_symlink(path, linkname, member, payload):
    import io
    with tarfile.open(path, "w:gz") as t:
        li = tarfile.TarInfo("link"); li.type = tarfile.SYMTYPE; li.linkname = linkname
        t.addfile(li)
        fi = tarfile.TarInfo(member); fi.size = len(payload)
        t.addfile(fi, io.BytesIO(payload))

def security_cases():
    """返回 [(名称, 执行函数)]，每个函数返回 (ok, detail)。"""
    def traversal():
        d = tempfile.mkdtemp(prefix="sec-trav-")
        try:
            victim = os.path.join(d, "victim.txt")
            open(victim, "w").write("ORIGINAL")
            work = os.path.join(d, "work"); os.makedirs(work)
            evil = os.path.join(d, "evil.tar.gz")
            _mk_tar_with_member(evil, "../victim.txt", b"PWNED")
            r = subprocess.run([sys.executable, ENGINE, evil, "zip", "-o", work],
                               capture_output=True, text=True, timeout=120)
            got = open(victim).read()
            ok = got == "ORIGINAL"
            return ok, "" if ok else f"tar 穿越覆盖了外部文件: victim={got!r} rc={r.returncode}"
        finally:
            shutil.rmtree(d, ignore_errors=True)

    def symlink_escape():
        d = tempfile.mkdtemp(prefix="sec-sym-")
        try:
            victim = os.path.join(d, "victim.txt")
            open(victim, "w").write("ORIGINAL")
            work = os.path.join(d, "work"); os.makedirs(work)
            evil = os.path.join(d, "evil.tar.gz")
            _mk_tar_with_symlink(evil, victim, "link", b"PWNED VIA SYMLINK")
            subprocess.run([sys.executable, ENGINE, evil, "zip", "-o", work],
                           capture_output=True, text=True, timeout=120)
            got = open(victim).read()
            ok = got == "ORIGINAL"
            return ok, "" if ok else f"tar symlink 逃逸覆盖了外部文件: victim={got!r}"
        finally:
            shutil.rmtree(d, ignore_errors=True)

    def no_overwrite():
        d = tempfile.mkdtemp(prefix="sec-ow-")
        try:
            from PIL import Image
            Image.new("RGB", (600, 600), (0, 0, 255)).save(os.path.join(d, "pic.jpg"), "JPEG")
            Image.new("RGB", (50, 50), (255, 0, 0)).save(os.path.join(d, "pic.png"))
            r = subprocess.run([sys.executable, ENGINE, os.path.join(d, "pic.png"), "jpg", "-o", d],
                               capture_output=True, text=True, timeout=120)
            payload = json.loads(r.stdout.strip().splitlines()[-1])
            old_size = Image.open(os.path.join(d, "pic.jpg")).size
            new_exists = os.path.exists(os.path.join(d, "pic-1.jpg"))
            ok = old_size == (600, 600) and new_exists and payload.get("ok")
            return ok, "" if ok else f"同名文件被覆盖: pic.jpg={old_size}, pic-1.jpg存在={new_exists}"
        finally:
            shutil.rmtree(d, ignore_errors=True)

    def force_overwrites():
        d = tempfile.mkdtemp(prefix="sec-force-")
        try:
            from PIL import Image
            Image.new("RGB", (600, 600), (0, 0, 255)).save(os.path.join(d, "pic.jpg"), "JPEG")
            Image.new("RGB", (50, 50), (255, 0, 0)).save(os.path.join(d, "pic.png"))
            subprocess.run([sys.executable, ENGINE, os.path.join(d, "pic.png"), "jpg", "-o", d, "--force"],
                           capture_output=True, text=True, timeout=120)
            ok = Image.open(os.path.join(d, "pic.jpg")).size == (50, 50)
            return ok, "" if ok else "--force 未按预期覆盖"
        finally:
            shutil.rmtree(d, ignore_errors=True)

    def archive_roundtrip():
        """7z 产出必须能被自己读回来（曾经因 '.' 条目必然失败）。"""
        d = tempfile.mkdtemp(prefix="sec-arc-")
        try:
            z = os.path.join(d, "pack.zip")
            with zipfile.ZipFile(z, "w") as zf:
                zf.writestr("a/b.txt", "hello")
            r1 = subprocess.run([sys.executable, ENGINE, z, "7z", "-o", d],
                                capture_output=True, text=True, timeout=120)
            p1 = json.loads(r1.stdout.strip().splitlines()[-1])
            if not p1.get("ok"): return False, f"zip->7z 失败: {p1.get('error')}"
            r2 = subprocess.run([sys.executable, ENGINE, p1["output"], "zip", "-o", os.path.join(d, "rt")],
                                capture_output=True, text=True, timeout=120)
            p2 = json.loads(r2.stdout.strip().splitlines()[-1])
            if not p2.get("ok"): return False, f"7z->zip 回环失败: {p2.get('error')}"
            with zipfile.ZipFile(p2["output"]) as zf:
                ok = "a/b.txt" in zf.namelist()
            return ok, "" if ok else f"回环后内容缺失: {zf.namelist()}"
        finally:
            shutil.rmtree(d, ignore_errors=True)

    def archive_name():
        d = tempfile.mkdtemp(prefix="sec-name-")
        try:
            src = "/tmp/nonexistent"
            z = os.path.join(d, "pack.tar.gz")
            with tarfile.open(z, "w:gz") as t:
                import io
                ti = tarfile.TarInfo("x.txt"); data = b"hi"; ti.size = len(data)
                t.addfile(ti, io.BytesIO(data))
            r = subprocess.run([sys.executable, ENGINE, z, "zip", "-o", os.path.join(d, "out")],
                               capture_output=True, text=True, timeout=120)
            p = json.loads(r.stdout.strip().splitlines()[-1])
            ok = p.get("ok") and os.path.basename(p["output"]) == "pack.zip"
            return ok, "" if ok else f"命名错误: {os.path.basename(p.get('output') or '')}"
        finally:
            shutil.rmtree(d, ignore_errors=True)

    return [("tar 路径穿越", traversal), ("tar symlink 逃逸", symlink_escape),
            ("同名不覆盖", no_overwrite), ("--force 覆盖", force_overwrites),
            ("7z 回环", archive_roundtrip), ("tar.gz 命名", archive_name)]

def main():
    shutil.rmtree(OUT, ignore_errors=True); os.makedirs(OUT)
    results_path = os.path.join(ROOT, "results.jsonl")
    passed, failed = 0, []
    with open(results_path, "w") as rf:
        for i, (fixture, tgt, validator) in enumerate(CASES):
            cid = f"case{i:03d}_{os.path.splitext(fixture)[0]}_{tgt.replace('.','_')}"
            cdir = os.path.join(OUT, cid); os.makedirs(cdir)
            src = os.path.join(FIX, fixture)
            r = subprocess.run([sys.executable, ENGINE, src, tgt, "-o", cdir], capture_output=True, text=True, timeout=600)
            status, detail = "FAIL", ""
            try:
                payload = json.loads(r.stdout.strip().splitlines()[-1]) if r.stdout.strip() else {}
                assert payload.get("ok"), f"engine not ok: {payload.get('error') or r.stderr[-200:]}"
                op = payload["output"]
                assert op and os.path.exists(op) and os.path.getsize(op) > 0, "output missing/empty"
                validator(op)
                status = "PASS"
            except Exception as e:
                detail = str(e)[:300]
            rec = {"case": cid, "src": fixture, "tgt": tgt, "status": status, "detail": detail}
            rf.write(json.dumps(rec, ensure_ascii=False) + "\n"); rf.flush()
            if status == "PASS": passed += 1
            else: failed.append(rec)
            print(f"[{i+1}/{len(CASES)}] {fixture} -> {tgt}: {status} {detail}", flush=True)
        for j, (fixture, tgt) in enumerate(NEGATIVES):
            cid = f"neg{j}_{os.path.splitext(fixture)[0]}_{tgt}"
            cdir = os.path.join(OUT, cid); os.makedirs(cdir)
            r = subprocess.run([sys.executable, ENGINE, os.path.join(FIX, fixture), tgt, "-o", cdir], capture_output=True, text=True, timeout=300)
            ok = r.returncode == 2 and not os.listdir(cdir)
            rec = {"case": cid, "src": fixture, "tgt": tgt, "status": "PASS" if ok else "FAIL",
                   "detail": "" if ok else f"expected rejection rc=2 no-file, got rc={r.returncode} files={os.listdir(cdir)} out={r.stdout[-150:]}"}
            rf.write(json.dumps(rec, ensure_ascii=False) + "\n"); rf.flush()
            if ok: passed += 1
            else: failed.append(rec)
            print(f"[neg{j+1}/{len(NEGATIVES)}] {fixture} -> {tgt}: {'PASS' if ok else 'FAIL'} {rec['detail']}", flush=True)

        sec = security_cases()
        for k, (name, fn) in enumerate(sec):
            cid = f"sec{k}_{name}"
            try:
                ok, detail = fn()
            except Exception as e:
                ok, detail = False, f"{type(e).__name__}: {e}"
            rec = {"case": cid, "src": name, "tgt": "-", "status": "PASS" if ok else "FAIL", "detail": detail}
            rf.write(json.dumps(rec, ensure_ascii=False) + "\n"); rf.flush()
            if ok: passed += 1
            else: failed.append(rec)
            print(f"[sec{k+1}/{len(sec)}] {name}: {'PASS' if ok else 'FAIL'} {detail}", flush=True)

    total = len(CASES) + len(NEGATIVES) + len(sec)
    summary = f"TOTAL {total} | PASS {passed} | FAIL {len(failed)}\n"
    for f in failed: summary += f"FAILED: {f['src']} -> {f['tgt']} :: {f['detail']}\n"
    with open(os.path.join(ROOT, "last_run.txt"), "w") as fh: fh.write(summary)
    print(summary, flush=True)

if __name__ == "__main__":
    main()
